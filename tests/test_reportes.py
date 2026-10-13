"""Reportes en Excel del coordinador: clientes, inscripciones, talleres y asistencia."""
from datetime import date
from io import BytesIO
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from app.models.dao.Reporte_DAO import ReporteDAO
from app.models.dto.Taller_DTO import TallerDTO

SERVICIO = "app.services.Reporte_Servicio."
TIPO_EXCEL = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILA_ENCABEZADO = 5


@pytest.fixture
def reporte_dao():
    with patch(SERVICIO + "ReporteDAO") as dao, patch(SERVICIO + "TallerDAO") as taller_dao:
        taller_dao.obtener_por_id.return_value = TallerDTO(5, "Liderazgo 1", 1, None, None, 1)
        dao.clientes_y_progreso.return_value = [{
            "cedula": "222", "nombres": "Ana", "apellidos": "Pérez", "correo": "ana@correo.co", "telefono": "300",
            "edad": 25, "estado": "Activo", "nivel_actual": "Avanzado",
            "inscripciones": 3, "aprobados": 1, "no_aprobados": 1, "programa_completado": 0,
        }]
        dao.inscripciones.return_value = [{
            "fecha_inscripcion": date(2026, 10, 1), "cedula": "222", "nombres": "Ana", "apellidos": "Pérez",
            "correo": "ana@correo.co", "taller": "Liderazgo 1", "nivel": "Basico",
            "fecha_inicio": date(2026, 10, 16), "fecha_fin": date(2026, 10, 18), "coach": "Carlos", "estado": "Iniciado",
        }]
        dao.talleres.return_value = [{
            "nombre": "Liderazgo 1", "nivel": "Basico", "fecha_inicio": date(2026, 10, 16), "fecha_fin": date(2026, 10, 18),
            "estado": "Activo", "coach": "Carlos", "coordinador": "Admin", "personal_capacitado": "Laura",
            "inscritos": 5, "aprobados": 3, "no_aprobados": 1,
        }]
        dao.asistencia.return_value = [{
            "taller": "Liderazgo 1", "fecha_inicio": date(2026, 10, 16), "fecha_fin": date(2026, 10, 18),
            "cedula": "222", "nombres": "Ana", "apellidos": "Pérez", "estado": "Iniciado", "sesiones": 4, "asistencias": 3,
        }]
        yield {"reporte": dao, "taller": taller_dao}


def _descargar(client, cabeceras, ruta):
    respuesta = client.get(ruta, headers=cabeceras)
    assert respuesta.status_code == 200, respuesta.json
    assert respuesta.mimetype == TIPO_EXCEL
    hoja = load_workbook(BytesIO(respuesta.data)).active
    encabezados = [celda.value for celda in hoja[FILA_ENCABEZADO]]
    primera_fila = dict(zip(encabezados, [celda.value for celda in hoja[FILA_ENCABEZADO + 1]]))
    return respuesta, hoja, primera_fila


class TestPermisos:

    def test_sin_token(self, client):
        assert client.get("/api/reportes/clientes").status_code == 401

    @pytest.mark.parametrize("rol", ["como_coach", "como_personal_capacitado", "como_cliente"])
    def test_solo_el_coordinador(self, client, request, rol):
        assert client.get("/api/reportes/clientes", headers=request.getfixturevalue(rol)).status_code == 403

    def test_reporte_inexistente(self, client, como_coordinador):
        assert client.get("/api/reportes/otro", headers=como_coordinador).status_code == 404


class TestContenido:

    def test_clientes_y_progreso(self, client, reporte_dao, como_coordinador):
        respuesta, hoja, fila = _descargar(client, como_coordinador, "/api/reportes/clientes")
        assert hoja["A1"].value == "Clientes y progreso"
        assert hoja["A2"].value.startswith("Lidera360 · Generado el ")
        assert fila["Cliente"] == "Ana Pérez"
        assert fila["En curso"] == 1
        assert fila["Programa completado"] == "No"
        assert "Lidera360-clientes-" in respuesta.headers["Content-Disposition"]

    def test_inscripciones(self, client, reporte_dao, como_coordinador):
        _, hoja, fila = _descargar(client, como_coordinador, "/api/reportes/inscripciones")
        assert fila["Cliente"] == "Ana Pérez"
        assert fila["Fecha de inscripción"].date() == date(2026, 10, 1)
        assert hoja.cell(row=FILA_ENCABEZADO + 1, column=1).number_format == "DD/MM/YYYY"
        assert hoja["A3"].value is None

    def test_talleres_calcula_la_tasa_de_aprobacion(self, client, reporte_dao, como_coordinador):
        _, _, fila = _descargar(client, como_coordinador, "/api/reportes/talleres")
        assert fila["% aprobación (de evaluados)"] == 0.75
        assert fila["Personal capacitado"] == "Laura"

    def test_asistencia_calcula_el_porcentaje(self, client, reporte_dao, como_coordinador):
        _, _, fila = _descargar(client, como_coordinador, "/api/reportes/asistencia")
        assert fila["Participante"] == "Ana Pérez"
        assert fila["% asistencia"] == 0.75

    def test_sin_sesiones_no_divide_por_cero(self, client, reporte_dao, como_coordinador):
        reporte_dao["reporte"].asistencia.return_value[0].update({"sesiones": 0, "asistencias": 0})
        _, _, fila = _descargar(client, como_coordinador, "/api/reportes/asistencia")
        assert fila["% asistencia"] is None

    def test_reporte_vacio(self, client, reporte_dao, como_coordinador):
        reporte_dao["reporte"].talleres.return_value = []
        _, hoja, _ = _descargar(client, como_coordinador, "/api/reportes/talleres")
        assert hoja.cell(row=FILA_ENCABEZADO + 1, column=1).value == "No hay datos para los filtros seleccionados."


class TestFiltros:

    def test_inscripciones_por_fechas_y_taller(self, client, reporte_dao, como_coordinador):
        _, hoja, _ = _descargar(
            client, como_coordinador,
            "/api/reportes/inscripciones?fecha_desde=2026-10-01&fecha_hasta=2026-10-31&id_taller=5",
        )
        reporte_dao["reporte"].inscripciones.assert_called_once_with(date(2026, 10, 1), date(2026, 10, 31), 5)
        assert hoja["A3"].value == "Filtros: Desde 01/10/2026 · Hasta 31/10/2026 · Taller: Liderazgo 1"

    def test_asistencia_por_taller(self, client, reporte_dao, como_coordinador):
        _descargar(client, como_coordinador, "/api/reportes/asistencia?id_taller=5")
        reporte_dao["reporte"].asistencia.assert_called_once_with(5)

    @pytest.mark.parametrize("consulta, mensaje", [
        ("fecha_desde=01-10-2026", "La fecha inicial debe tener el formato AAAA-MM-DD"),
        ("fecha_desde=2026-10-31&fecha_hasta=2026-10-01", "La fecha final no puede ser anterior a la fecha inicial"),
        ("id_taller=abc", "El taller seleccionado no es válido"),
    ], ids=["fecha-mal-formada", "rango-invertido", "taller-no-numerico"])
    def test_valida_los_filtros(self, client, reporte_dao, como_coordinador, consulta, mensaje):
        respuesta = client.get(f"/api/reportes/inscripciones?{consulta}", headers=como_coordinador)
        assert respuesta.status_code == 400
        assert respuesta.json["mensaje"] == mensaje

    def test_taller_inexistente(self, client, reporte_dao, como_coordinador):
        reporte_dao["taller"].obtener_por_id.return_value = None
        assert client.get("/api/reportes/asistencia?id_taller=99", headers=como_coordinador).status_code == 404


class TestConsultasDAO:

    def test_inscripciones_usa_parametros_y_no_concatena_valores(self):
        with patch("app.models.dao.Reporte_DAO._consultar") as consultar:
            ReporteDAO.inscripciones(date(2026, 10, 1), date(2026, 10, 31), 5)
        sql, parametros = consultar.call_args[0]
        assert "i.fecha_inscripcion >= %s" in sql and "i.fecha_inscripcion <= %s" in sql and "i.id_taller = %s" in sql
        assert parametros == (date(2026, 10, 1), date(2026, 10, 31), 5)

    def test_inscripciones_sin_filtros(self):
        with patch("app.models.dao.Reporte_DAO._consultar") as consultar:
            ReporteDAO.inscripciones()
        sql, parametros = consultar.call_args[0]
        assert "WHERE" not in sql.split("FROM inscripcion i")[1]
        assert parametros == ()


def test_el_navegador_puede_leer_el_nombre_del_archivo(client, reporte_dao, como_coordinador):
    respuesta = client.get("/api/reportes/clientes", headers={**como_coordinador, "Origin": "https://frontend.pruebas"})
    assert "Content-Disposition" in respuesta.headers.get("Access-Control-Expose-Headers", "")
