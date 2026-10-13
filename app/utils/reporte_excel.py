from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.utils.fechas import ahora

MORADO = "7A3CE0"
MORADO_OSCURO = "2B123F"
LILA = "82639F"

FORMATOS = {
    "fecha": "DD/MM/YYYY",
    "fecha_hora": "DD/MM/YYYY HH:MM",
    "porcentaje": "0%",
    "entero": "0",
}


def generar_excel(titulo, columnas, filas, filtros=None):

    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]

    hoja["A1"] = titulo
    hoja["A1"].font = Font(bold=True, size=14, color=MORADO_OSCURO)
    hoja["A2"] = f"Lidera360 · Generado el {ahora().strftime('%d/%m/%Y %H:%M')}"
    hoja["A2"].font = Font(size=10, color=LILA)
    if filtros:
        hoja["A3"] = "Filtros: " + " · ".join(filtros)
        hoja["A3"].font = Font(size=10, color=LILA)

    fila_encabezado = 5
    for indice, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_encabezado, column=indice, value=columna["titulo"])
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=MORADO)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(indice)].width = columna.get("ancho", 16)

    for numero, fila in enumerate(filas, start=fila_encabezado + 1):
        for indice, columna in enumerate(columnas, start=1):
            valor = fila.get(columna["clave"])
            if isinstance(valor, datetime) and columna.get("formato") != "fecha_hora":
                valor = valor.date()
            celda = hoja.cell(row=numero, column=indice, value=valor)
            formato = FORMATOS.get(columna.get("formato"))
            if formato and valor is not None:
                celda.number_format = formato
            if columna.get("ajustar"):
                celda.alignment = Alignment(wrap_text=True, vertical="top")

    ultima_fila = fila_encabezado + max(len(filas), 1)
    hoja.auto_filter.ref = f"A{fila_encabezado}:{get_column_letter(len(columnas))}{ultima_fila}"
    hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)

    if not filas:
        hoja.cell(row=fila_encabezado + 1, column=1, value="No hay datos para los filtros seleccionados.")

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
